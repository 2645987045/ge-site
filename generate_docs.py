# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_DIR = r"D:\community"
FONT_NAME = "SimSun"  # 宋体
FONT_NAME_ASCII = "SimSun"

def set_run_font(run, size=Pt(10), bold=False, color=None, italic=False):
    """统一设置 run 的字体为宋体，避免字体不一致。"""
    run.font.name = FONT_NAME
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    # 设置 East Asian 字体
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), FONT_NAME_ASCII)
    rFonts.set(qn('w:hAnsi'), FONT_NAME_ASCII)
    rFonts.set(qn('w:eastAsia'), FONT_NAME)
    rFonts.set(qn('w:cs'), FONT_NAME_ASCII)

def set_paragraph_font(paragraph, size=Pt(10), bold=False, color=None):
    """设置段落中所有 run 的字体。"""
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)

def set_cell_shading(cell, color):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def set_default_doc_font(doc):
    """设置文档默认字体为宋体。"""
    style = doc.styles['Normal']
    font = style.font
    font.name = FONT_NAME
    font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), FONT_NAME_ASCII)
    rFonts.set(qn('w:hAnsi'), FONT_NAME_ASCII)
    rFonts.set(qn('w:eastAsia'), FONT_NAME)
    rFonts.set(qn('w:cs'), FONT_NAME_ASCII)
    
    # 设置所有标题样式的字体
    for i in range(1, 5):
        style_name = f'Heading {i}'
        if style_name in doc.styles:
            hs = doc.styles[style_name]
            hs.font.name = FONT_NAME
            rpr = hs.element.get_or_add_rPr()
            rFonts = rpr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rpr.insert(0, rFonts)
            rFonts.set(qn('w:ascii'), FONT_NAME_ASCII)
            rFonts.set(qn('w:hAnsi'), FONT_NAME_ASCII)
            rFonts.set(qn('w:eastAsia'), FONT_NAME)
            rFonts.set(qn('w:cs'), FONT_NAME_ASCII)
    
    # 设置 List Bullet 样式
    if 'List Bullet' in doc.styles:
        ls = doc.styles['List Bullet']
        ls.font.name = FONT_NAME
        rpr = ls.element.get_or_add_rPr()
        rFonts = rpr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rpr.insert(0, rFonts)
        rFonts.set(qn('w:ascii'), FONT_NAME_ASCII)
        rFonts.set(qn('w:hAnsi'), FONT_NAME_ASCII)
        rFonts.set(qn('w:eastAsia'), FONT_NAME)

def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=Pt(16 if level == 1 else 13), bold=True, color=RGBColor(101, 67, 33))
    return h

def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row - bold, white, 宋体
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=Pt(10), bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(cell, '8B4513')
    
    # Data rows - normal weight, 宋体
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            set_run_font(run, size=Pt(9), bold=False)
            if r_idx % 2 == 1:
                set_cell_shading(cell, 'F5F0EB')
    
    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Cm(width)
    
    return table

def add_body_text(doc, text, bold=False, italic=False, size=Pt(10), color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p

def add_bullet_item(doc, label, text):
    """添加带标签的列表项，标签加粗，正文不加粗，统一宋体。"""
    p = doc.add_paragraph(style='List Bullet')
    # 清除默认 run
    p.clear()
    if label:
        run_label = p.add_run(label)
        set_run_font(run_label, size=Pt(10), bold=True)
    run_text = p.add_run(text)
    set_run_font(run_text, size=Pt(10), bold=False)
    return p

def add_bullet_plain(doc, text):
    """添加纯文本列表项，不加粗，统一宋体。"""
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=Pt(10), bold=False)
    return p


# ============================================================
# Document 1: 策划方案
# ============================================================
def create_plan_doc():
    doc = Document()
    set_default_doc_font(doc)
    
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('《吴玉章在高师》')
    set_run_font(run, size=Pt(26), bold=True, color=RGBColor(101, 67, 33))
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('票根文创策划方案')
    set_run_font(run, size=Pt(18), bold=True, color=RGBColor(139, 69, 19))
    
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run('项目类型：学生话剧纪念文创  |  风格：复古做旧  |  日期：2026年8月22日')
    set_run_font(run, size=Pt(9), color=RGBColor(128, 128, 128))
    
    doc.add_page_break()
    
    # ---- Section 1 ----
    add_heading_styled(doc, '一、项目概述', level=1)
    rows = [
        ['话剧名称', '《吴玉章在高师》'],
        ['话剧类型', '学生话剧，严肃庄重正剧风格'],
        ['主题内容', '吴玉章先生生平事迹（成都高等师范学校时期）'],
        ['文创品类', '纪念票根（第一期）'],
        ['票根定位', '纪念收藏品，非实际入场券'],
        ['核心创意', '复古实体票根 + 数字互动网站（二维码联动）'],
    ]
    add_styled_table(doc, ['项目', '内容'], rows, col_widths=[4, 12])
    doc.add_paragraph()
    
    # ---- Section 2 ----
    add_heading_styled(doc, '二、票根设计规格', level=1)
    add_heading_styled(doc, '2.1 物理规格', level=2)
    rows = [
        ['尺寸', '210mm x 70mm（标准演出票根尺寸）'],
        ['材质', '牛皮纸（建议 120g-150g）'],
        ['印刷', '复古做旧风格，单色或双色印刷（深棕/暗红 + 黑色）'],
        ['工艺', '哑光覆膜（可选），增加手感与耐久性'],
        ['裁切', '左侧预留虚线撕裂线，模拟真实票根体验'],
    ]
    add_styled_table(doc, ['参数', '规格'], rows, col_widths=[4, 12])
    doc.add_paragraph()
    
    add_heading_styled(doc, '2.2 视觉风格', level=2)
    items = [
        ('整体调性：', '复古做旧，呼应民国/早期高等教育历史氛围'),
        ('色彩方案：', '以牛皮纸原色为底，搭配深棕、暗红、墨黑等沉稳色调'),
        ('字体建议：', '标题使用衬线体/书法体（如方正清刻本悦宋），正文使用楷体或仿宋'),
        ('装饰元素：', '做旧边框、印章纹样、老票据底纹、细线分隔等传统票据元素'),
        ('差异化亮点：', '票根融合数字互动（二维码），实现"复古外壳 + 现代内核"'),
    ]
    for label, text in items:
        add_bullet_item(doc, label, text)
    doc.add_paragraph()
    
    # ---- Section 3 ----
    add_heading_styled(doc, '三、票根内容规划', level=1)
    add_heading_styled(doc, '3.1 正面（主视觉面）', level=2)
    for item in [
        '剧名《吴玉章在高师》（突出展示）',
        'Slogan / 宣传语（后期确定后加入）',
        '主视觉图形（Logo、吴玉章先生剪影、或高师建筑元素等）',
        '演出信息（时间、地点、场次——印"首演"信息增强仪式感）',
        '二维码（引导语："扫码探索话剧世界"或"扫码进入高师记忆"）',
    ]:
        add_bullet_plain(doc, item)
    
    add_heading_styled(doc, '3.2 背面（内容面）', level=2)
    for item in [
        '收藏编号（如 No.001、No.002，限量编号增加收藏价值）',
        '经典台词节选（1-2句，从剧中精选）',
        '吴玉章先生名言（1句）',
        '剧情梗概或创作背景（2-3行简短文字）',
        '制作团队信息（出品方、监制、设计等）',
    ]:
        add_bullet_plain(doc, item)
    doc.add_paragraph()
    
    # ---- Section 4 ----
    add_heading_styled(doc, '四、配套网站规划', level=1)
    add_heading_styled(doc, '4.1 网站入口', level=2)
    add_body_text(doc, '票根正面二维码为唯一入口，建议网站域名简洁易记。')
    
    add_heading_styled(doc, '4.2 首期内容模块', level=2)
    rows = [
        ['介绍视频', '话剧宣传片/预告片，建议 2-5 分钟'],
        ['角色档案', '剧中主要角色介绍、饰演者信息、角色历史原型'],
        ['幕后花絮', '排练照片、幕后故事、主创访谈'],
    ]
    add_styled_table(doc, ['模块', '内容说明'], rows, col_widths=[4, 12])
    doc.add_paragraph()
    
    add_heading_styled(doc, '4.3 后续可扩展模块', level=2)
    for item in [
        '吴玉章先生生平史料专区',
        '观众留言墙 / 互动区',
        '隐藏彩蛋（未公开剧照、导演手记等）',
        '其他文创产品展示与购买入口',
    ]:
        add_bullet_plain(doc, item)
    
    add_heading_styled(doc, '4.4 网站风格建议', level=2)
    for item in [
        '与票根的复古做旧风格保持统一',
        '配色以暖色调为主，搭配历史感纹理背景',
        '响应式设计，适配手机扫码后浏览',
    ]:
        add_bullet_plain(doc, item)
    doc.add_paragraph()
    
    # ---- Section 5 ----
    add_heading_styled(doc, '五、执行排期', level=1)
    rows = [
        ['第一阶段', '内容准备', '确定 slogan、精选台词/名言、撰写梗概', '1 周'],
        ['第二阶段', '视觉设计', '票根正反面设计稿（含 2-3 版方案对比）', '1-2 周'],
        ['第三阶段', '网站搭建', '网站基础框架搭建 + 首期内容上线', '2-3 周'],
        ['第四阶段', '打样确认', '票根印刷打样、色彩校对、材质确认', '1 周'],
        ['第五阶段', '批量印刷', '确认终稿后批量印刷', '1-2 周'],
        ['第六阶段', '发放使用', '演出现场发放 / 搭配其他文创组合发放', '演出日'],
    ]
    add_styled_table(doc, ['阶段', '任务', '内容', '周期'], rows, col_widths=[2.5, 2.5, 8, 3])
    doc.add_paragraph()
    
    # ---- Section 6 ----
    add_heading_styled(doc, '六、成本估算参考', level=1)
    rows = [
        ['票根设计', '0 元（AI 辅助设计）', '如需专业设计师另计'],
        ['网站搭建', '0-500 元', '使用免费建站工具或学生优惠'],
        ['域名', '30-80 元/年', '视域名后缀而定'],
        ['印刷（200张）', '100-300 元', '牛皮纸 + 单/双色印刷'],
        ['印刷（500张）', '200-500 元', '量大单价更低'],
    ]
    add_styled_table(doc, ['项目', '预估费用', '备注'], rows, col_widths=[4, 4, 8])
    
    p = doc.add_paragraph()
    run = p.add_run('省钱建议：')
    set_run_font(run, size=Pt(10), bold=True)
    run = p.add_run('可使用线上印刷平台（如"世纪开元""虎彩"等），上传设计稿直接下单，小批量起印价格友好。')
    set_run_font(run, size=Pt(10), bold=False)
    doc.add_paragraph()
    
    # ---- Section 7 ----
    add_heading_styled(doc, '七、后续文创扩展方向（第二期）', level=1)
    items = [
        ('主题笔记本 / 手账本：', '封面印剧目海报，内页印台词节选'),
        ('场刊 / 演出手册：', '完整剧情介绍 + 主创团队 + 幕后故事'),
        ('明信片套装：', '剧照 / 插画 / 名言'),
        ('金属书签：', '吴玉章先生剪影 / 高师建筑轮廓'),
        ('徽章 / 胸针：', '剧目 Logo / 经典元素'),
        ('帆布袋：', 'slogan + 主视觉'),
    ]
    for label, text in items:
        add_bullet_item(doc, label, text)
    
    path = os.path.join(OUTPUT_DIR, '吴玉章在高师-票根文创策划方案.docx')
    doc.save(path)
    print(f'策划方案已保存: {path}')
    return path


# ============================================================
# Document 2: 详尽任务表
# ============================================================
def create_task_doc():
    doc = Document()
    set_default_doc_font(doc)
    
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('《吴玉章在高师》票根文创')
    set_run_font(run, size=Pt(24), bold=True, color=RGBColor(101, 67, 33))
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('执行任务表')
    set_run_font(run, size=Pt(18), bold=True, color=RGBColor(139, 69, 19))
    
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run('创建日期：2026年8月22日  |  总工期预估：6-10 周  |  状态标记：[ ] 待办  [~] 进行中  [x] 完成')
    set_run_font(run, size=Pt(9), color=RGBColor(128, 128, 128))
    
    doc.add_page_break()
    
    # Helper to add phase
    def add_phase(title, goal, tasks):
        add_heading_styled(doc, title, level=1)
        p = doc.add_paragraph()
        run = p.add_run(goal)
        set_run_font(run, size=Pt(10), italic=True)
        add_styled_table(doc, ['编号', '状态', '任务名称', '具体说明', '负责人', '截止时间', '前置依赖'], tasks, col_widths=[1.2, 1, 3, 5.5, 1.8, 1.5, 2.5])
        doc.add_paragraph()
    
    # Phase 1
    add_phase('第一阶段：内容准备（第1周）',
        '目标：确定所有文字内容素材，为设计阶段提供完整文案。',
        [
            ['1.1', '[ ]', '确定 Slogan / 宣传语', '根据话剧主题创作 5-10 组候选宣传语，团队投票选定最终版', '文案组', '第1周', '剧本、吴玉章史料'],
            ['1.2', '[ ]', '精选经典台词', '从剧本中筛选 3-5 句最具代表性的台词，用于票根背面', '文案组/导演', '第1周', '完整剧本'],
            ['1.3', '[ ]', '整理吴玉章名言', '收集吴玉章先生经典语录，选取 1-2 句适合票根的', '文案组', '第1周', '吴玉章文集/史料'],
            ['1.4', '[ ]', '撰写剧情梗概', '撰写 2-3 行票根背面简介（100字以内）', '文案组', '第1周', '剧本'],
            ['1.5', '[ ]', '确定演出信息', '确认首演时间、地点、场次等具体信息', '制片组', '第1周', '演出排期表'],
            ['1.6', '[ ]', '整理制作团队名单', '汇总出品方、监制、导演、设计等人员信息', '制片组', '第1周', '团队花名册'],
            ['1.7', '[ ]', '确定收藏编号规则', '决定编号方式（按场次/日期/流水号），确定首批总量', '制片组', '第1周', '—'],
        ])
    
    # Phase 2
    add_phase('第二阶段：视觉设计（第2-3周）',
        '目标：完成票根正反面设计定稿，确定主视觉风格。',
        [
            ['2.1', '[ ]', '收集参考素材', '收集民国老票据、复古印刷品等视觉参考，建立情绪板', '设计组', '第2周', '—'],
            ['2.2', '[ ]', '设计主视觉元素', '创作票根主视觉：Logo/吴玉章剪影/高师建筑元素（至少2版）', '设计组', '第2周', '第一阶段文案'],
            ['2.3', '[ ]', '票根正面设计稿', '完成正面排版设计（剧名、主视觉、演出信息、二维码区域）', '设计组', '第2-3周', '2.1, 2.2'],
            ['2.4', '[ ]', '票根背面设计稿', '完成背面排版设计（编号、台词、名言、梗概、团队信息）', '设计组', '第2-3周', '第一阶段文案'],
            ['2.5', '[ ]', '出 2-3 版方案对比', '制作 2-3 套不同风格方案供团队评审', '设计组', '第3周', '2.3, 2.4'],
            ['2.6', '[ ]', '团队评审定稿', '团队内部投票/讨论，确定最终设计方案', '全组', '第3周', '2.5'],
            ['2.7', '[ ]', '设计稿终稿输出', '输出印刷用高清文件（CMYK、300dpi、含出血位）', '设计组', '第3周', '2.6'],
            ['2.8', '[ ]', '生成二维码', '生成指向网站的二维码，嵌入票根设计', '技术组', '第3周', '网站域名确定'],
        ])
    
    # Phase 3
    add_phase('第三阶段：配套网站搭建（第3-5周）',
        '目标：搭建票根扫码落地网站，首期内容上线。',
        [
            ['3.1', '[ ]', '确定技术方案', '选择建站方式（静态站/WordPress/Notion站点等），确定域名', '技术组', '第3周', '—'],
            ['3.2', '[ ]', '注册域名与部署', '购买域名、配置服务器/托管平台（GitHub Pages/Vercel等）', '技术组', '第3周', '3.1'],
            ['3.3', '[ ]', '网站 UI 设计', '设计与票根风格统一的网站视觉（复古做旧风）', '设计组', '第3-4周', '2.6 定稿风格'],
            ['3.4', '[ ]', '制作介绍视频', '剪辑话剧宣传片/预告片（2-5分钟）', '视频组', '第3-4周', '排练素材'],
            ['3.5', '[ ]', '编写角色档案', '整理各角色的介绍、饰演者、历史原型信息', '文案组', '第3-4周', '剧本'],
            ['3.6', '[ ]', '整理幕后花絮', '收集排练照片、幕后故事、主创访谈素材', '宣传组', '第3-4周', '排练期间素材'],
            ['3.7', '[ ]', '网站开发搭建', '实现网站页面（首页+视频+角色+花絮模块）', '技术组', '第4-5周', '3.2, 3.3'],
            ['3.8', '[ ]', '内容填充上线', '将视频、角色档案、花絮等内容上传至网站', '技术组/文案组', '第5周', '3.7, 3.4-3.6'],
            ['3.9', '[ ]', '移动端测试', '测试手机扫码访问体验，确保响应式布局正常', '全组', '第5周', '3.8'],
            ['3.10', '[ ]', '网站正式上线', '确认无误后正式上线，二维码可正常访问', '技术组', '第5周', '3.9'],
        ])
    
    # Phase 4
    add_phase('第四阶段：打样确认（第6周）',
        '目标：确认印刷品质，确保材质、色彩、工艺符合预期。',
        [
            ['4.1', '[ ]', '选择印刷供应商', '对比 2-3 家印刷平台（世纪开元/虎彩/本地印刷厂），确认报价', '制片组', '第6周', '—'],
            ['4.2', '[ ]', '提交打样文件', '将终稿设计文件提交印刷方，确认格式要求（CMYK/出血位等）', '设计组', '第6周', '2.7'],
            ['4.3', '[ ]', '收到打样并评审', '检查打样实物：色彩偏差、材质手感、撕裂线效果、二维码清晰度', '全组', '第6周', '4.2'],
            ['4.4', '[ ]', '色彩校对与修正', '如打样有色差，调整设计文件后重新打样', '设计组', '第6周', '4.3'],
            ['4.5', '[ ]', '确认最终样品', '签字确认最终样品，作为批量印刷标准', '制片组', '第6周', '4.4'],
        ])
    
    # Phase 5
    add_phase('第五阶段：批量印刷（第7-8周）',
        '目标：完成批量印刷，质检入库。',
        [
            ['5.1', '[ ]', '确定印刷数量', '根据演出场次、发放计划确定最终印刷数量', '制片组', '第7周', '—'],
            ['5.2', '[ ]', '下单批量印刷', '向确认的供应商下单，交付终稿文件', '制片组', '第7周', '4.5'],
            ['5.3', '[ ]', '跟进印刷进度', '与印刷方保持沟通，确认交期', '制片组', '第7-8周', '5.2'],
            ['5.4', '[ ]', '收货质检', '收到成品后抽检：色彩、裁切、覆膜质量、二维码可扫率', '制片组', '第8周', '5.3'],
            ['5.5', '[ ]', '入库管理', '按编号整理入库，记录库存数量', '制片组', '第8周', '5.4'],
        ])
    
    # Phase 6
    add_phase('第六阶段：发放使用（演出日）',
        '目标：在演出现场有序发放/售卖纪念票根。',
        [
            ['6.1', '[ ]', '制定发放方案', '确定发放方式（购票赠送/单独售卖/限量编号发放）', '制片组', '演出前1周', '—'],
            ['6.2', '[ ]', '准备发放物料', '准备发放用的收纳盒/信封、说明卡片等', '制片组', '演出前1周', '—'],
            ['6.3', '[ ]', '安排现场人员', '安排票根发放/售卖的现场工作人员', '制片组', '演出前3天', '6.1'],
            ['6.4', '[ ]', '现场发放/售卖', '演出当天在签到处/出口处发放或售卖', '现场组', '演出日', '5.5, 6.3'],
            ['6.5', '[ ]', '收集观众反馈', '记录观众对票根的评价和建议，为后续文创参考', '宣传组', '演出日', '6.4'],
        ])
    
    doc.add_page_break()
    
    # 附一
    add_heading_styled(doc, '附一：建议人员分工', level=1)
    rows = [
        ['文案组', 'Slogan创作、台词精选、名言整理、剧情梗概、角色档案文案', '2-3人'],
        ['设计组', '主视觉设计、票根正反面排版、网站UI、印刷文件输出', '1-2人'],
        ['技术组', '网站搭建、域名配置、二维码生成、移动端测试', '1-2人'],
        ['视频组', '宣传片拍摄/剪辑、幕后花絮视频制作', '1-2人'],
        ['宣传组', '幕后素材收集、观众反馈收集、社交媒体宣传', '1-2人'],
        ['制片组', '排期管理、供应商对接、印刷跟进、现场发放', '1-2人'],
    ]
    add_styled_table(doc, ['角色', '职责范围', '建议人数'], rows, col_widths=[2.5, 9, 2.5])
    doc.add_paragraph()
    
    # 附二
    add_heading_styled(doc, '附二：关键里程碑', level=1)
    rows = [
        ['M1', '文案定稿', '所有文字内容（slogan、台词、名言、梗概）确认完毕', '第1周末'],
        ['M2', '设计定稿', '票根正反面设计终稿通过评审', '第3周末'],
        ['M3', '网站上线', '配套网站正式上线，二维码可访问', '第5周末'],
        ['M4', '打样通过', '印刷打样确认，色彩材质达标', '第6周末'],
        ['M5', '成品入库', '批量印刷完成，质检入库', '第8周末'],
        ['M6', '正式发放', '演出现场发放/售卖', '演出日'],
    ]
    add_styled_table(doc, ['节点', '里程碑', '达成标准', '目标时间'], rows, col_widths=[1.5, 3, 7, 3])
    doc.add_paragraph()
    
    # 附三
    add_heading_styled(doc, '附三：风险与注意事项', level=1)
    rows = [
        ['印刷色差', '牛皮纸吸墨性强，实际色彩可能与屏幕显示有偏差', '务必打样确认，预留修正时间'],
        ['二维码失效', '网站域名过期或服务器故障导致二维码无法访问', '选择可靠托管平台，设置域名自动续费'],
        ['工期延误', '设计反复修改或印刷排期紧张', '每个阶段预留 2-3 天缓冲期'],
        ['版权风险', '使用的字体、图片需确认授权', '优先使用免费商用字体和原创素材'],
        ['预算超支', '特殊工艺或加急印刷会增加成本', '提前确认报价，避免临时加急'],
    ]
    add_styled_table(doc, ['风险项', '说明', '应对措施'], rows, col_widths=[2.5, 6, 6.5])
    doc.add_paragraph()
    
    # 附四
    add_heading_styled(doc, '附四：AI 可辅助完成的任务', level=1)
    p = doc.add_paragraph()
    run = p.add_run('以下任务可使用 AI 工具辅助完成，提高效率：')
    set_run_font(run, size=Pt(10))
    
    rows = [
        ['Slogan 创作', '根据话剧主题、吴玉章精神生成多组宣传语候选', '高'],
        ['台词精选建议', '分析剧本，推荐最具代表性和感染力的台词', '高'],
        ['名言整理', '汇总吴玉章先生经典语录并筛选适合票根的', '高'],
        ['剧情梗概撰写', '根据剧本生成 2-3 行精炼简介', '高'],
        ['票根视觉设计', '生成复古做旧风格的票根设计稿/参考图', '中'],
        ['网站 UI 设计', '生成与票根风格统一的网站页面设计', '中'],
        ['二维码生成', '生成指向网站的二维码图片', '高'],
        ['编号规则设计', '设计收藏编号体系和生成编号序列', '中'],
        ['角色档案撰写', '根据剧本和历史资料生成角色介绍文案', '高'],
    ]
    add_styled_table(doc, ['任务', 'AI 辅助方式', '可行性'], rows, col_widths=[3, 8, 2])
    
    path = os.path.join(OUTPUT_DIR, '吴玉章在高师-票根文创任务表.docx')
    doc.save(path)
    print(f'任务表已保存: {path}')
    return path


if __name__ == '__main__':
    create_plan_doc()
    create_task_doc()
    print('两份文档均已生成完毕！')
